"""Generate MGM Status Report PowerPoint slide."""

import os

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import nsmap
from datetime import datetime

def RgbColor(r, g, b):
    """Create RGB color for pptx."""
    from pptx.dml.color import RGBColor
    return RGBColor(r, g, b)

def create_mgm_status_pptx():
    """Create MGM status report PowerPoint."""
    
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Slide 1: Executive Summary
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)
    
    # Title bar
    title_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.9))
    title_shape.fill.solid()
    title_shape.fill.fore_color.rgb = RgbColor(4, 159, 217)  # Cisco blue
    title_shape.line.fill.background()
    
    title_box = slide.shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(10), Inches(0.6))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "MGM Resorts - Delivery Status Report"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RgbColor(255, 255, 255)
    
    # Date
    date_box = slide.shapes.add_textbox(Inches(10.5), Inches(0.25), Inches(2.5), Inches(0.5))
    tf = date_box.text_frame
    p = tf.paragraphs[0]
    p.text = "April 10, 2026"
    p.font.size = Pt(14)
    p.font.color.rgb = RgbColor(255, 255, 255)
    p.alignment = PP_ALIGN.RIGHT
    
    # Status indicator
    status_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11), Inches(1.1), Inches(2), Inches(0.5))
    status_box.fill.solid()
    status_box.fill.fore_color.rgb = RgbColor(255, 193, 7)  # Yellow
    status_box.line.fill.background()
    
    status_text = slide.shapes.add_textbox(Inches(11), Inches(1.15), Inches(2), Inches(0.4))
    tf = status_text.text_frame
    p = tf.paragraphs[0]
    p.text = "🟡 YELLOW"
    p.font.size = Pt(16)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    
    # Executive Summary section
    summary_title = slide.shapes.add_textbox(Inches(0.3), Inches(1.1), Inches(5), Inches(0.4))
    tf = summary_title.text_frame
    p = tf.paragraphs[0]
    p.text = "Executive Summary"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RgbColor(30, 58, 95)
    
    summary_content = slide.shapes.add_textbox(Inches(0.3), Inches(1.5), Inches(6.2), Inches(2))
    tf = summary_content.text_frame
    tf.word_wrap = True
    
    summary_text = """• Firewall migration PO pending MGM internal review
• Technical discovery kickoff scheduled for today (Apr 10)
• Plan buildout target: 1-1.5 weeks
• CRITICAL: Palo Alto support expires April 15th
• ISE automation (ISAAC) progressing - minor PSN05 issue
• Strong MGM partnership - PAN "scorched earth" offer declined"""
    
    p = tf.paragraphs[0]
    p.text = summary_text
    p.font.size = Pt(12)
    p.line_spacing = 1.3
    
    # Key Milestones & Target Dates
    milestones_title = slide.shapes.add_textbox(Inches(6.8), Inches(1.1), Inches(6), Inches(0.4))
    tf = milestones_title.text_frame
    p = tf.paragraphs[0]
    p.text = "Key Milestones & Target Dates"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RgbColor(30, 58, 95)
    
    # Milestones table
    milestones = [
        ("Apr 10", "Technical Discovery Kickoff", "🟢"),
        ("Apr 11-12", "Complete Network/SCC Access", "🟡"),
        ("Apr 15", "⚠️ PAN Support Expires", "🔴"),
        ("Apr 17-18", "Migration Plan v1 Complete", "🟡"),
        ("Apr 21", "Netscout TAPs Call (Scott/Phil)", "🟡"),
        ("Apr 25", "Resource Allocation Finalized", "🟡"),
        ("May 1", "Phase 1 Migration Start (Target)", "⚪"),
    ]
    
    y_pos = 1.5
    for date, milestone, status in milestones:
        row = slide.shapes.add_textbox(Inches(6.8), Inches(y_pos), Inches(6), Inches(0.3))
        tf = row.text_frame
        p = tf.paragraphs[0]
        p.text = f"{status} {date}: {milestone}"
        p.font.size = Pt(11)
        if "⚠️" in milestone or status == "🔴":
            p.font.bold = True
            p.font.color.rgb = RgbColor(220, 53, 69)
        y_pos += 0.32
    
    # Recommended Next Steps section
    next_steps_title = slide.shapes.add_textbox(Inches(0.3), Inches(3.7), Inches(6), Inches(0.4))
    tf = next_steps_title.text_frame
    p = tf.paragraphs[0]
    p.text = "Recommended Next Steps"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RgbColor(30, 58, 95)
    
    next_steps = [
        ("IMMEDIATE", "Validate SCC access for Daniel/Mike Culp", "Today"),
        ("IMMEDIATE", "Engage MGM security team with Nexar", "Today"),
        ("HIGH", "Finalize PAN support gap mitigation plan", "Apr 14"),
        ("HIGH", "Increase Mike Culp allocation to 75-80%", "Apr 11"),
        ("HIGH", "Schedule Netscout TAPs call (Scott/Deepak/Phil)", "Apr 14"),
        ("MEDIUM", "Complete ISE PSN05 investigation", "Apr 12"),
        ("MEDIUM", "Identify Nexar 3rd party support group", "Apr 11"),
    ]
    
    y_pos = 4.1
    for priority, action, target in next_steps:
        # Priority badge
        if priority == "IMMEDIATE":
            color = RgbColor(220, 53, 69)  # Red
        elif priority == "HIGH":
            color = RgbColor(255, 193, 7)  # Yellow
        else:
            color = RgbColor(40, 167, 69)  # Green
        
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.3), Inches(y_pos), Inches(0.9), Inches(0.25))
        badge.fill.solid()
        badge.fill.fore_color.rgb = color
        badge.line.fill.background()
        
        badge_text = slide.shapes.add_textbox(Inches(0.3), Inches(y_pos + 0.02), Inches(0.9), Inches(0.22))
        tf = badge_text.text_frame
        p = tf.paragraphs[0]
        p.text = priority
        p.font.size = Pt(8)
        p.font.bold = True
        p.font.color.rgb = RgbColor(255, 255, 255) if priority != "HIGH" else RgbColor(0, 0, 0)
        p.alignment = PP_ALIGN.CENTER
        
        # Action text
        action_text = slide.shapes.add_textbox(Inches(1.3), Inches(y_pos), Inches(4), Inches(0.25))
        tf = action_text.text_frame
        p = tf.paragraphs[0]
        p.text = action
        p.font.size = Pt(10)
        
        # Target date
        target_text = slide.shapes.add_textbox(Inches(5.3), Inches(y_pos), Inches(1), Inches(0.25))
        tf = target_text.text_frame
        p = tf.paragraphs[0]
        p.text = target
        p.font.size = Pt(10)
        p.font.italic = True
        p.font.color.rgb = RgbColor(100, 100, 100)
        
        y_pos += 0.35
    
    # Risk Register
    risk_title = slide.shapes.add_textbox(Inches(6.8), Inches(3.7), Inches(6), Inches(0.4))
    tf = risk_title.text_frame
    p = tf.paragraphs[0]
    p.text = "Risk Register"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RgbColor(30, 58, 95)
    
    risks = [
        ("HIGH", "PAN support expires Apr 15 - no 3rd party coverage"),
        ("HIGH", "S2S VPN migration complexity (Mike Culp flagged)"),
        ("MED", "Resource constraints - Culp needs 75-80% allocation"),
        ("LOW", "ISE PSN05 certificate error - investigation ongoing"),
    ]
    
    y_pos = 4.1
    for severity, risk in risks:
        if severity == "HIGH":
            color = RgbColor(220, 53, 69)
        elif severity == "MED":
            color = RgbColor(255, 193, 7)
        else:
            color = RgbColor(40, 167, 69)
        
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(y_pos), Inches(0.6), Inches(0.25))
        badge.fill.solid()
        badge.fill.fore_color.rgb = color
        badge.line.fill.background()
        
        badge_text = slide.shapes.add_textbox(Inches(6.8), Inches(y_pos + 0.02), Inches(0.6), Inches(0.22))
        tf = badge_text.text_frame
        p = tf.paragraphs[0]
        p.text = severity
        p.font.size = Pt(8)
        p.font.bold = True
        p.font.color.rgb = RgbColor(255, 255, 255) if severity != "MED" else RgbColor(0, 0, 0)
        p.alignment = PP_ALIGN.CENTER
        
        risk_text = slide.shapes.add_textbox(Inches(7.5), Inches(y_pos), Inches(5.5), Inches(0.3))
        tf = risk_text.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = risk
        p.font.size = Pt(10)
        
        y_pos += 0.38
    
    # Active Workstreams
    workstreams_title = slide.shapes.add_textbox(Inches(0.3), Inches(6.3), Inches(12), Inches(0.4))
    tf = workstreams_title.text_frame
    p = tf.paragraphs[0]
    p.text = "Active Workstreams"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RgbColor(30, 58, 95)
    
    workstreams = [
        ("Firepower Migration", "Discovery phase", "🟡"),
        ("ISE Automation (ISAAC)", "Cert renewal in progress", "🟢"),
        ("LCS R&S Support", "Monitoring Cat 3K syslog", "🟢"),
        ("Secure Access", "Pending FW completion", "⚪"),
    ]
    
    x_pos = 0.3
    for name, status, indicator in workstreams:
        ws_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x_pos), Inches(6.7), Inches(3), Inches(0.55))
        ws_box.fill.solid()
        ws_box.fill.fore_color.rgb = RgbColor(241, 245, 249)
        ws_box.line.color.rgb = RgbColor(203, 213, 225)
        
        ws_text = slide.shapes.add_textbox(Inches(x_pos + 0.1), Inches(6.75), Inches(2.8), Inches(0.45))
        tf = ws_text.text_frame
        p = tf.paragraphs[0]
        p.text = f"{indicator} {name}"
        p.font.size = Pt(11)
        p.font.bold = True
        
        p2 = tf.add_paragraph()
        p2.text = status
        p2.font.size = Pt(9)
        p2.font.color.rgb = RgbColor(100, 100, 100)
        
        x_pos += 3.25
    
    # Footer
    footer = slide.shapes.add_textbox(Inches(0.3), Inches(7.2), Inches(12), Inches(0.3))
    tf = footer.text_frame
    p = tf.paragraphs[0]
    p.text = "Source: Webex Spacelift Export Analysis | Key Contacts: Mike Brown (Delivery Lead), Jason Anderson (CX Leadership), Paul Snow (Sales)"
    p.font.size = Pt(8)
    p.font.color.rgb = RgbColor(150, 150, 150)
    
    # Save
    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "MGM_Status_Report_20260410.pptx")
    prs.save(output_path)
    print(f"PowerPoint saved to: {output_path}")
    return output_path

if __name__ == "__main__":
    create_mgm_status_pptx()
