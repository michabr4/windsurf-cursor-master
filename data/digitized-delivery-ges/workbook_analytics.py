"""
Post-processing for NaC readiness workbook: parity-gap defaults and an Excel table
on the main grid. Pivot-style summaries are added by native_interactive_pivots (static tables).
"""

from __future__ import annotations

from typing import Any

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table, TableStyleInfo


PRIORITY_TECHS = frozenset({"SD-WAN", "Cat-C/SDA"})
GAP_COMMENT = "Parity exercise not completed"


def _last_data_row(ws: Any, c_customer: int) -> int:
    last = 1
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=c_customer).value
        if v is not None and str(v).strip():
            last = r
    return last


def _pct_missing(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and not value.strip():
        return True
    return False


def _has_parity_survey_trace(comments: Any) -> bool:
    s = str(comments or "")
    return (
        "Parity (raw):" in s
        or "Report: http" in s
        or "Report: https" in s
    )


def apply_sdwan_catcsda_parity_gaps(ws: Any, h: dict[str, int]) -> int:
    """
    For SD-WAN / Cat-C/SDA rows with no parity % and no merged survey trace in comments,
    set NaC Parity Complete = No and ensure NaC Parity Comments includes the gap note.
    Returns count of rows updated.
    """
    c_tech = h["Technologies"]
    c_pct = h["NaC Parity %"]
    c_complete = h["NaC Parity Complete"]
    c_comments = h["NaC Parity Comments"]
    c_customer = h["Customer Name"]
    updated = 0
    for r in range(2, ws.max_row + 1):
        if not ws.cell(row=r, column=c_customer).value or not str(ws.cell(row=r, column=c_customer).value).strip():
            continue
        tech = ws.cell(row=r, column=c_tech).value
        if tech is None or str(tech).strip() not in PRIORITY_TECHS:
            continue
        pct = ws.cell(row=r, column=c_pct).value
        if not _pct_missing(pct):
            continue
        comments = ws.cell(row=r, column=c_comments).value
        if _has_parity_survey_trace(comments):
            continue

        ws.cell(row=r, column=c_complete, value="No")
        cm = str(comments).strip() if comments else ""
        if not cm:
            ws.cell(row=r, column=c_comments, value=GAP_COMMENT)
        elif GAP_COMMENT not in cm:
            ws.cell(row=r, column=c_comments, value=f"{cm} | {GAP_COMMENT}")
        updated += 1
    return updated


def fill_missing_nac_score_formulas(ws: Any, h: dict[str, int]) -> int:
    """
    Copy the NaC Score array formula onto rows that have a customer but no formula.

    New rows appended by the merge only set parity columns; Excel does not auto-fill
    NaC Score. Template rows use MainTable[[#This Row], ...] — the same formula text
    is valid on every table row; only the per-cell array ref changes.
    """
    c_score = h.get("NaC Score")
    c_customer = h.get("Customer Name")
    if c_score is None or c_customer is None:
        return 0

    template_text: str | None = None
    for r in range(2, min(ws.max_row + 1, 5000)):
        v = ws.cell(row=r, column=c_score).value
        if isinstance(v, ArrayFormula) and v.text:
            template_text = v.text
            break
    if not template_text:
        return 0

    col_letter = get_column_letter(c_score)
    last_r = _last_data_row(ws, c_customer)
    filled = 0
    for r in range(2, last_r + 1):
        cust = ws.cell(row=r, column=c_customer).value
        if cust is None or not str(cust).strip():
            continue
        cell = ws.cell(row=r, column=c_score)
        if cell.data_type == "f":
            continue
        cell.value = ArrayFormula(ref=f"{col_letter}{r}", text=template_text)
        filled += 1
    return filled


def add_data_table_to_score_sheet(ws: Any, h: dict[str, int]) -> None:
    """Structured table on `as-code score` (MainTable) for filtering and optional user pivots."""
    c_customer = h["Customer Name"]
    last_r = _last_data_row(ws, c_customer)
    if last_r < 2:
        return
    max_c = max(h.values())
    ref = f"A1:{get_column_letter(max_c)}{last_r}"
    # Must stay "MainTable": NaC Score (and other) formulas use structured refs like MainTable[[#This Row],...].
    display_name = "MainTable"
    for nm in list(ws.tables):
        del ws.tables[nm]
    tab = Table(displayName=display_name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(tab)
