#!/usr/bin/env python3
"""
Export a workbook to a new file.

Default: copy .xlsx → new .xlsx (preserves formulas, formatting, and sheets).
Optional: --xls writes legacy .xls (values only; dates as strings).
"""

from __future__ import annotations

import argparse
import shutil
from datetime import date, datetime, time
from pathlib import Path


def export_xlsx_copy(src: Path, dst: Path) -> None:
    shutil.copy2(src, dst)


def export_xls_values(src: Path, dst: Path) -> None:
    import openpyxl
    import xlwt

    wb_in = openpyxl.load_workbook(src, read_only=True, data_only=True)
    wb_out = xlwt.Workbook()
    for sheet_name in wb_in.sheetnames:
        ws_in = wb_in[sheet_name]
        ws_out = wb_out.add_sheet(sheet_name[:31])
        for r_idx, row in enumerate(ws_in.iter_rows(values_only=True)):
            for c_idx, val in enumerate(row):
                if val is None:
                    continue
                if isinstance(val, (datetime, date, time)):
                    ws_out.write(r_idx, c_idx, str(val))
                elif isinstance(val, (int, float)):
                    ws_out.write(r_idx, c_idx, val)
                else:
                    ws_out.write(r_idx, c_idx, str(val))
    wb_in.close()
    wb_out.save(str(dst))


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument(
        "src",
        nargs="?",
        default="SaC-NaC Readiness Checklist v1_merged.xlsx",
        type=Path,
        help="Source .xlsx path",
    )
    p.add_argument(
        "dst",
        nargs="?",
        default=None,
        type=Path,
        help="Destination path (default: <src stem>_export.xlsx or .xls with --xls)",
    )
    p.add_argument(
        "--xls",
        action="store_true",
        help="Write legacy .xls instead of copying .xlsx",
    )
    args = p.parse_args()

    if args.dst:
        dst = args.dst
    elif args.xls:
        dst = args.src.with_name(args.src.stem + "_export.xls")
    else:
        dst = args.src.with_name(args.src.stem + "_export.xlsx")

    dst.parent.mkdir(parents=True, exist_ok=True)

    if args.xls:
        export_xls_values(args.src, dst)
    else:
        export_xlsx_copy(args.src, dst)

    print(f"Wrote {dst.resolve()}")


if __name__ == "__main__":
    main()
