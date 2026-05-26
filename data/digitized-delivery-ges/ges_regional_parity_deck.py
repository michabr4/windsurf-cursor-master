#!/usr/bin/env python3
"""
Build a slim regional PowerPoint: one slide per GES customer segment, SD-WAN + Cat-C/SDA only.

Each regional slide lists four buckets in the body placeholder (outline levels), split by
technology — same “Title, 1 Column with Bullets” layout as phase2_customer_criteria_deck.py.

Copies SaC_NaC_Adoption_Plan.pptx (default) for slide masters, theme fonts, and layouts.
"""

from __future__ import annotations

import argparse
import re
import shutil
from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Pt

TECH_ORDER = ("SD-WAN", "Cat-C/SDA")
PHASE1_TECHS = frozenset(TECH_ORDER)
NO_ACCESS_PHRASES = (
    "n/a",
    "na",
    "no access",
    "no direct access",
    "don't have direct access",
    "dont have direct access",
    "waiting to get access",
    "error in collection",
    "no network",
    "without network",
    "lack of access",
    "unable to access",
    "cannot access",
)

REGION_SORT = {
    "GES-East": 0,
    "GES-West": 1,
    "GES-Central": 2,
    "GES-Premier": 3,
    "Other": 4,
}

# Template layout indices (SaC_NaC_Adoption_Plan.pptx) — match Phase 2 criteria deck.
LAYOUT_TITLE_SUBTITLE = 10  # Title, Subtitle Only 1
LAYOUT_TITLE_1COL_BULLETS = 19  # Title, 1 Column with Bullets


def _col_index(headers: list[Any], name: str) -> int:
    for i, h in enumerate(headers, start=1):
        if h == name:
            return i
    raise KeyError(name)


def _is_num(v: Any) -> bool:
    return isinstance(v, (int, float)) and v == v


def _pct_value(v: Any) -> float | None:
    if not _is_num(v):
        return None
    return float(v)


def _norm_text(s: Any) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def _complete_yes(v: Any) -> bool:
    s = str(v or "").strip().lower()
    return s in ("yes", "y", "true", "1")


def _no_access_context(comments: Any, net_access: Any) -> bool:
    blob = _norm_text(comments) + " " + _norm_text(net_access)
    return any(p in blob for p in NO_ACCESS_PHRASES)


@dataclass(frozen=True)
class ScopedRow:
    customer: str
    technology: str
    segment: str
    parity_pct: Any
    parity_complete: Any
    comments: str
    network_access: str


def load_scoped_rows(xlsx: Path) -> list[ScopedRow]:
    wb = load_workbook(xlsx, data_only=True)
    ws = wb["as-code score"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    c_cust = _col_index(headers, "Customer Name")
    c_tech = _col_index(headers, "Technologies")
    c_pct = _col_index(headers, "NaC Parity %")
    c_comp = _col_index(headers, "NaC Parity Complete")
    c_seg = _col_index(headers, "Customer Segment")
    c_comments = _col_index(headers, "NaC Parity Comments")
    c_net = _col_index(headers, "Customer Network Access")
    out: list[ScopedRow] = []
    for r in range(2, ws.max_row + 1):
        cust = ws.cell(r, c_cust).value
        if cust is None or not str(cust).strip():
            continue
        tech = str(ws.cell(r, c_tech).value or "").strip()
        if tech not in PHASE1_TECHS:
            continue
        seg_v = ws.cell(r, c_seg).value
        seg = str(seg_v).strip() if seg_v is not None and str(seg_v).strip() else ""
        out.append(
            ScopedRow(
                customer=str(cust).strip(),
                technology=tech,
                segment=seg,
                parity_pct=ws.cell(r, c_pct).value,
                parity_complete=ws.cell(r, c_comp).value,
                comments=str(ws.cell(r, c_comments).value or ""),
                network_access=str(ws.cell(r, c_net).value or ""),
            )
        )
    wb.close()
    return out


def classify_row(r: ScopedRow) -> int:
    """Return bucket 1–4 for this scoped row."""
    pct = _pct_value(r.parity_pct)
    if pct is not None and pct >= 0.80:
        return 1
    if (
        _complete_yes(r.parity_complete)
        and (pct is None or pct == 0.0)
        and _no_access_context(r.comments, r.network_access)
    ):
        return 2
    if pct is not None and 0.0 < pct < 0.80:
        return 3
    return 4


def _segment_label(seg: str) -> str:
    return seg if seg else "Unassigned"


def _region_sort_key(seg: str) -> tuple[int, str]:
    label = _segment_label(seg)
    return (REGION_SORT.get(label, 50), label.casefold())


def bucket_sets_by_region(rows: list[ScopedRow]) -> dict[str, dict[int, dict[str, set[str]]]]:
    """
    region_label -> bucket (1..4) -> tech -> set of customer names
    """
    acc: dict[str, dict[int, dict[str, set[str]]]] = defaultdict(
        lambda: defaultdict(lambda: defaultdict(set))
    )
    for r in rows:
        seg_label = _segment_label(r.segment)
        b = classify_row(r)
        acc[seg_label][b][r.technology].add(r.customer)
    return acc


def _body_shape(slide: Any) -> Any:
    for shape in slide.shapes:
        if shape.is_placeholder and shape.placeholder_format.type == PP_PLACEHOLDER.BODY:
            return shape
    raise RuntimeError("No BODY placeholder on slide")


def _fill_body_single_column(
    slide: Any,
    paragraphs: list[tuple[int, str]],
    *,
    level0_pt: int = 11,
    level1_pt: int = 11,
) -> None:
    """paragraphs: list of (outline_level, text). level0 = section / tech lines; level1 = customers."""
    body = _body_shape(slide)
    tf = body.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, (level, text) in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.level = level
        p.font.size = Pt(level0_pt if level == 0 else level1_pt)
        p.font.bold = level == 0
        p.space_after = Pt(2)


def _paragraphs_for_region(
    buckets: dict[int, dict[str, set[str]]],
    names_cap: int,
) -> list[tuple[int, str]]:
    labels = {
        1: "1) Customers at ≥80% parity (by technology)",
        2: "2) Parity complete @ 0% — no network access (NaC Parity Comments + Customer Network Access)",
        3: "3) Customers with parity >0% and <80%",
        4: "4) Parity not done for this technology (no numeric %, incomplete, or outside 1–3)",
    }
    lines: list[tuple[int, str]] = []
    for bid in (1, 2, 3, 4):
        lines.append((0, labels[bid]))
        tech_map = buckets.get(bid, {})
        for tech in TECH_ORDER:
            names = sorted(tech_map.get(tech, set()), key=str.casefold)
            lines.append((0, f"{tech} ({len(names)})"))
            if not names:
                lines.append((1, "—"))
            else:
                for n in names[:names_cap]:
                    lines.append((1, n))
                if len(names) > names_cap:
                    lines.append((1, f"+ {len(names) - names_cap} more (see workbook)"))
    return lines


def _strip_all_slides(prs: Presentation) -> None:
    while len(prs.slides) > 0:
        rId = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[0]


def add_title_slide(prs: Presentation, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TITLE_SUBTITLE])
    slide.shapes.title.text = "NaC Parity — GES regional view"
    ttp = slide.shapes.title.text_frame.paragraphs[0]
    ttp.font.size = Pt(40)
    ttp.font.bold = True
    body = _body_shape(slide)
    tf = body.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = subtitle
    p.level = 0
    p.font.bold = False
    p.font.size = Pt(20)


def add_region_slide(
    prs: Presentation,
    region: str,
    buckets: dict[int, dict[str, set[str]]],
    names_cap: int,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TITLE_1COL_BULLETS])
    slide.shapes.title.text = f"{region} — Cat-C/SDA & SD-WAN"
    rtp = slide.shapes.title.text_frame.paragraphs[0]
    rtp.font.size = Pt(32)
    rtp.font.bold = True
    _fill_body_single_column(
        slide,
        _paragraphs_for_region(buckets, names_cap),
        level0_pt=16,
        level1_pt=14,
    )


def build_presentation(
    *,
    base_ppt: Path,
    output_ppt: Path,
    rows: list[ScopedRow],
    names_cap: int,
) -> int:
    shutil.copy2(base_ppt, output_ppt)
    prs = Presentation(str(output_ppt))
    _strip_all_slides(prs)

    add_title_slide(
        prs,
        f"SD-WAN & Cat-C/SDA · Source: merged readiness workbook · {date.today().isoformat()}",
    )

    acc = bucket_sets_by_region(rows)
    regions = sorted(acc.keys(), key=_region_sort_key)
    for reg in regions:
        add_region_slide(prs, reg, acc[reg], names_cap=names_cap)

    output_ppt.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_ppt))
    return len(prs.slides)


def main() -> None:
    ap = argparse.ArgumentParser(description="GES regional NaC parity deck (corporate template).")
    ap.add_argument(
        "--base-ppt",
        type=Path,
        default=Path("SaC_NaC_Adoption_Plan.pptx"),
        help="Template deck for theme/fonts/layouts (same as Phase 2 criteria deck)",
    )
    ap.add_argument(
        "--xlsx",
        type=Path,
        default=Path("SaC-NaC Readiness Checklist v1_merged.xlsx"),
        help="Merged readiness workbook",
    )
    ap.add_argument(
        "-o",
        "--output",
        type=Path,
        default=Path("GES_Regional_Parity_Status.pptx"),
        help="Output PowerPoint path",
    )
    ap.add_argument(
        "--names-cap",
        type=int,
        default=38,
        help="Max customer lines per technology per section before truncation (Phase 2 default: 38)",
    )
    args = ap.parse_args()

    rows = load_scoped_rows(args.xlsx)
    n = build_presentation(
        base_ppt=args.base_ppt,
        output_ppt=args.output,
        rows=rows,
        names_cap=args.names_cap,
    )
    print(f"Wrote {args.output.resolve()} ({n} slides)")


if __name__ == "__main__":
    main()
